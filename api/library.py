"""
Library: per-user documents & artifacts (issue #14).

Every user has their own library (users/<uid>/library/, metadata in the DB):
uploads, documents OpenMuse generates, filled PDF forms, and (later) browser
downloads and email attachments. The agent works with them through docs.*:

  docs.list        (R0) what's in the library
  docs.read        (R1) text of a document (PDF text layer, md/txt/csv, xlsx cells)
  docs.pdf_fields  (R1) fillable fields of a PDF form
  docs.pdf_fill    (R2) fill a form -> a NEW filled copy (the original is never
                        modified); signature fields are left for the user
  docs.create      (R2) generate Markdown / CSV / XLSX / PDF (PDF rendered by
                        headless Chromium from Markdown)

Safety: text-bearing files are secret-scanned before they're stored; PDFs
that carry JavaScript/launch actions are rejected; document text reaches the
model as untrusted data (like every tool result).
"""
from __future__ import annotations

import base64
import hashlib
import html
import io
import json
import os
import re
import threading
import time
import uuid

from tools.redaction import looks_like_secret

MAX_BYTES = 15 * 1024 * 1024
MIME = {".pdf": "application/pdf", ".md": "text/markdown", ".txt": "text/plain", ".csv": "text/csv",
        ".json": "application/json", ".html": "text/html",
        ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".gif": "image/gif",
        ".webp": "image/webp"}
_PDF_ACTIVE = re.compile(rb"/(JavaScript|JS|Launch|EmbeddedFile|RichMedia)\b")


class Library:
    def __init__(self, backend, root: str):
        self.backend = backend
        self.root = root
        self._mem: dict[str, dict] = {}
        self._lock = threading.RLock()

    # -- storage --------------------------------------------------------------
    def _dir(self, user_id: str) -> str:
        d = os.path.join(self.root, re.sub(r"[^A-Za-z0-9_-]", "_", user_id or "user_api"), "library")
        os.makedirs(d, exist_ok=True)
        return d

    def _put(self, rec: dict) -> None:
        if self.backend.db is not None:
            self.backend.db.kv_put("library", rec["artifact_id"], rec, user_id=rec["user_id"])
        else:
            self._mem[rec["artifact_id"]] = rec

    def list(self, user_id: str) -> list[dict]:
        if self.backend.db is not None:
            items = self.backend.db.kv_list("library", user_id=user_id, limit=1000)
        else:
            items = [r for r in self._mem.values() if r["user_id"] == user_id]
        return sorted(items, key=lambda r: r["created_at"], reverse=True)

    def get(self, user_id: str, artifact_id: str) -> dict | None:
        rec = (self.backend.db.kv_get("library", artifact_id) if self.backend.db is not None
               else self._mem.get(artifact_id))
        return rec if rec and rec["user_id"] == user_id else None

    def data(self, rec: dict) -> bytes:
        with open(os.path.join(self._dir(rec["user_id"]), rec["file"]), "rb") as fh:
            return fh.read()

    def add(self, user_id: str, name: str, data: bytes, *, source: str = "upload",
            meta: dict | None = None) -> dict:
        if not data:
            raise ValueError("empty file")
        if len(data) > MAX_BYTES:
            raise ValueError("file too large (15 MB max)")
        name = os.path.basename(name or "document").strip() or "document"
        ext = os.path.splitext(name)[1].lower()
        mime = MIME.get(ext, "application/octet-stream")
        if ext == ".pdf":
            if not data.startswith(b"%PDF"):
                raise ValueError("not a valid PDF")
            if _PDF_ACTIVE.search(data):
                raise ValueError("this PDF contains active content (JavaScript/launch actions) and was rejected")
        if mime.startswith("text/") or ext in (".json", ".md", ".csv"):
            if looks_like_secret(data[:200_000].decode("utf-8", "replace")):
                raise ValueError("this file looks like it contains secrets; not stored")
        aid = "art_" + uuid.uuid4().hex[:12]
        fname = aid + (ext or ".bin")
        with open(os.path.join(self._dir(user_id), fname), "wb") as fh:
            fh.write(data)
        rec = {"artifact_id": aid, "user_id": user_id, "name": name, "mime": mime, "size": len(data),
               "sha256": hashlib.sha256(data).hexdigest(), "source": source, "file": fname,
               "created_at": time.time(), "meta": meta or {}}
        self._put(rec)
        return rec

    def delete(self, user_id: str, artifact_id: str) -> bool:
        rec = self.get(user_id, artifact_id)
        if rec is None:
            return False
        try:
            os.remove(os.path.join(self._dir(user_id), rec["file"]))
        except OSError:
            pass
        if self.backend.db is not None:
            self.backend.db.kv_delete("library", artifact_id)
        else:
            self._mem.pop(artifact_id, None)
        return True

    @staticmethod
    def public(rec: dict) -> dict:
        return {k: rec[k] for k in ("artifact_id", "name", "mime", "size", "source", "created_at", "meta")}

    # -- reading ------------------------------------------------------------------
    def text(self, rec: dict, limit: int = 20_000) -> str:
        data = self.data(rec)
        if rec["mime"] == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            out = "\n\n".join((p.extract_text() or "") for p in reader.pages)
        elif rec["mime"].endswith("spreadsheetml.sheet"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                rows.append(f"# {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    rows.append("\t".join("" if v is None else str(v) for v in row))
            out = "\n".join(rows)
        elif rec["mime"].startswith("text/") or rec["mime"] == "application/json":
            out = data.decode("utf-8", "replace")
        else:
            out = f"[{rec['mime']} file, {rec['size']} bytes — no text layer]"
        return out[:limit]

    # -- PDF forms --------------------------------------------------------------------
    def pdf_fields(self, rec: dict) -> list[dict]:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(self.data(rec)))
        fields = reader.get_fields() or {}
        out = []
        for name, f in fields.items():
            ftype = str(f.get("/FT", ""))
            kind = {"/Tx": "text", "/Btn": "checkbox", "/Ch": "choice", "/Sig": "signature"}.get(ftype, ftype or "text")
            opts = [str(o[1] if isinstance(o, list) else o) for o in (f.get("/Opt") or [])][:20]
            states = []
            if kind == "checkbox" and f.get("/_States_"):
                states = [str(s) for s in f["/_States_"]]
            out.append({"name": name, "type": kind, "label": str(f.get("/TU") or name),
                        "value": str(f.get("/V") or ""), "options": opts or states})
        return out

    def pdf_fill(self, user_id: str, rec: dict, values: dict) -> tuple[dict, list[dict]]:
        from pypdf import PdfReader, PdfWriter
        fields = {f["name"]: f for f in self.pdf_fields(rec)}
        unknown = [k for k in values if k not in fields]
        if unknown:
            raise ValueError(f"unknown field(s): {', '.join(unknown[:5])}")
        sigs = [k for k in values if fields[k]["type"] == "signature"]
        if sigs:
            raise ValueError("signature fields are left for the user to sign")
        from pypdf.generic import NameObject
        reader = PdfReader(io.BytesIO(self.data(rec)))
        writer = PdfWriter()
        writer.append(reader)
        text_vals = {k: str(v) for k, v in values.items() if fields[k]["type"] != "checkbox"}
        box_vals = {k: str(v).lower() in ("true", "yes", "1", "on", "x")
                    for k, v in values.items() if fields[k]["type"] == "checkbox"}
        for page in writer.pages:
            if text_vals:
                writer.update_page_form_field_values(page, text_vals, auto_regenerate=False)
            # checkboxes: set /V and /AS on the widget directly, using the form's
            # own "on" state name when it defines appearances (else /Yes)
            for annot_ref in page.get("/Annots") or []:
                annot = annot_ref.get_object()
                name = str(annot.get("/T", ""))
                if name not in box_vals:
                    continue
                on = "/Yes"
                ap = annot.get("/AP")
                if ap is not None:
                    states = [str(k) for k in ap.get_object().get("/N", {}).keys() if str(k) != "/Off"]
                    on = states[0] if states else on
                state = NameObject(on if box_vals[name] else "/Off")
                annot[NameObject("/V")] = state
                annot[NameObject("/AS")] = state
        writer.set_need_appearances_writer(True)
        buf = io.BytesIO()
        writer.write(buf)
        base, _ = os.path.splitext(rec["name"])
        filled = self.add(user_id, f"{base} (filled).pdf", buf.getvalue(), source="filled form",
                          meta={"filled_from": rec["artifact_id"]})
        review = [{"label": fields[k]["label"],
                   "value": ("Checked" if box_vals[k] else "Unchecked") if k in box_vals else str(values[k])}
                  for k in values]
        return filled, review

    # -- generation ---------------------------------------------------------------------
    def create(self, user_id: str, *, title: str, fmt: str, content: str = "",
               columns: list | None = None, rows: list | None = None) -> dict:
        title = (title or "Document").strip()[:120]
        safe = re.sub(r"[^\w\- ]+", "", title).strip() or "document"
        if looks_like_secret(content or "") or looks_like_secret(json.dumps(rows or [])[:50_000]):
            raise ValueError("refusing to write secrets into a document")
        if fmt == "md":
            return self.add(user_id, safe + ".md", (content or "").encode("utf-8"), source="generated")
        if fmt == "csv":
            import csv
            buf = io.StringIO()
            w = csv.writer(buf)
            if columns:
                w.writerow(columns)
            for r in rows or []:
                w.writerow(r if isinstance(r, list) else [r])
            return self.add(user_id, safe + ".csv", buf.getvalue().encode("utf-8"), source="generated")
        if fmt == "xlsx":
            import openpyxl
            from openpyxl.styles import Font
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = safe[:31] or "Sheet1"
            if columns:
                ws.append([str(c) for c in columns])
                for cell in ws[1]:
                    cell.font = Font(bold=True)
            for r in rows or []:
                ws.append(r if isinstance(r, list) else [r])
            for col in ws.columns:
                width = max((len(str(c.value or "")) for c in col), default=8)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
            buf = io.BytesIO()
            wb.save(buf)
            return self.add(user_id, safe + ".xlsx", buf.getvalue(), source="generated")
        if fmt == "pdf":
            return self.add(user_id, safe + ".pdf", render_pdf(title, content or ""), source="generated")
        raise ValueError("format must be md, csv, xlsx or pdf")


# -- Markdown -> HTML -> PDF (headless Chromium) ------------------------------------------
def markdown_html(title: str, md: str) -> str:
    out, in_list, table = [], False, []

    def inline(t):
        t = html.escape(t)
        t = re.sub(r"\*\*(.+?)\*\*", r"<strong>\1</strong>", t)
        t = re.sub(r"(?<![*\w])\*(?!\s)(.+?)\*(?!\w)", r"<em>\1</em>", t)
        t = re.sub(r"`([^`]+)`", r"<code>\1</code>", t)
        return re.sub(r"\[([^\]]+)\]\((https?://[^)\s]+)\)", r'<a href="\2">\1</a>', t)

    def flush_table():
        if not table:
            return
        rows = [[c.strip() for c in r.strip().strip("|").split("|")] for r in table]
        head, body = rows[0], [r for r in rows[1:] if not re.fullmatch(r"[\s:-]*", "".join(r))]
        out.append("<table><thead><tr>" + "".join(f"<th>{inline(c)}</th>" for c in head) + "</tr></thead><tbody>"
                   + "".join("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r) + "</tr>" for r in body)
                   + "</tbody></table>")
        table.clear()

    for line in (md or "").splitlines():
        if re.match(r"^\s*\|.*\|\s*$", line):
            table.append(line)
            continue
        flush_table()
        li = re.match(r"^\s*(?:[-*•]|\d+\.)\s+(.*)$", line)
        if li:
            if not in_list:
                out.append("<ul>"); in_list = True
            out.append(f"<li>{inline(li.group(1))}</li>")
            continue
        if in_list:
            out.append("</ul>"); in_list = False
        h = re.match(r"^(#{1,4})\s+(.*)$", line)
        if h:
            lvl = min(4, len(h.group(1)) + 1)
            out.append(f"<h{lvl}>{inline(h.group(2))}</h{lvl}>")
        elif line.strip():
            out.append(f"<p>{inline(line)}</p>")
    flush_table()
    if in_list:
        out.append("</ul>")
    css = ("body{font:12pt/1.5 -apple-system,Segoe UI,Helvetica,Arial,sans-serif;color:#16181d;margin:0}"
           "h1{font-size:20pt;margin:0 0 12pt}h2{font-size:15pt;margin:16pt 0 6pt}h3,h4{font-size:12.5pt}"
           "table{border-collapse:collapse;width:100%;margin:8pt 0}th,td{border-bottom:1px solid #ddd;"
           "padding:4pt 6pt;text-align:left}th{font-weight:600}code{font-family:Menlo,monospace;font-size:10pt}")
    return (f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
            f"<style>{css}</style></head><body><h1>{html.escape(title)}</h1>{''.join(out)}</body></html>")


def render_pdf(title: str, md: str) -> bytes:
    """Render Markdown to PDF with Playwright's Chromium (own thread; no network)."""
    result: dict = {}

    def work():
        try:
            from playwright.sync_api import sync_playwright
            with sync_playwright() as p:
                browser = p.chromium.launch()
                ctx = browser.new_context(java_script_enabled=False)
                page = ctx.new_page()
                page.route("**/*", lambda r: r.abort() if not r.request.url.startswith("data:") else r.continue_())
                page.set_content(markdown_html(title, md), wait_until="load")
                result["pdf"] = page.pdf(format="Letter", margin={"top": "0.8in", "bottom": "0.8in",
                                                                 "left": "0.8in", "right": "0.8in"})
                browser.close()
        except Exception as exc:
            result["error"] = exc
    t = threading.Thread(target=work, daemon=True)
    t.start()
    t.join(60)
    if "pdf" not in result:
        raise ValueError(f"PDF rendering failed: {result.get('error', 'timed out')}")
    return result["pdf"]


# -- agent tools ----------------------------------------------------------------------------
def register_tools(registry, library: Library) -> None:
    from tools.registry import ToolDefinition
    registry.register_namespace("docs", "The user's Library: read documents, fill PDF forms, create documents.")
    uid = lambda ctx: getattr(ctx, "user_id", "") or "user_api"

    def _rec(ctx, aid):
        rec = library.get(uid(ctx), aid)
        if rec is None:
            raise ValueError(f"no document {aid} in this user's Library (use docs.list)")
        return rec

    def card(rec, subtitle=""):
        return {"type": "document", "title": rec["name"], "artifact_id": rec["artifact_id"],
                "subtitle": subtitle or rec["source"], "mime": rec["mime"]}

    registry.register(ToolDefinition(
        name="docs.list", version="1.0.0", description="List documents in the user's Library.",
        input_schema={"type": "object", "properties": {}, "additionalProperties": False},
        output_schema={"type": "object"}, capabilities=["docs.read"], side_effect="none",
        idempotency="pure", default_timeout_ms=5_000,
        execute=lambda ctx, a: {"documents": [Library.public(r) for r in library.list(uid(ctx))[:50]]},
    ))
    registry.register(ToolDefinition(
        name="docs.read", version="1.0.0",
        description="Read the text of a Library document (PDF text layer, Markdown, CSV, text, spreadsheets).",
        input_schema={"type": "object", "properties": {"artifact_id": {"type": "string", "maxLength": 40}},
                      "required": ["artifact_id"], "additionalProperties": False},
        output_schema={"type": "object"}, capabilities=["docs.read"], side_effect="none",
        idempotency="pure", default_timeout_ms=30_000,
        execute=lambda ctx, a: {"name": _rec(ctx, a["artifact_id"])["name"],
                                "text": library.text(_rec(ctx, a["artifact_id"]))},
    ))
    registry.register(ToolDefinition(
        name="docs.pdf_fields", version="1.0.0",
        description="List the fillable fields of a PDF form in the Library (name, type, label, current value, options).",
        input_schema={"type": "object", "properties": {"artifact_id": {"type": "string", "maxLength": 40}},
                      "required": ["artifact_id"], "additionalProperties": False},
        output_schema={"type": "object"}, capabilities=["docs.read"], side_effect="none",
        idempotency="pure", default_timeout_ms=30_000,
        execute=lambda ctx, a: {"fields": library.pdf_fields(_rec(ctx, a["artifact_id"]))},
    ))

    def fill(ctx, a):
        filled, review = library.pdf_fill(uid(ctx), _rec(ctx, a["artifact_id"]), a["values"])
        return {"artifact_id": filled["artifact_id"], "name": filled["name"], "filled": review,
                "note": "Saved as a new copy; the original is unchanged. Signature fields are left for the user."}

    registry.register(ToolDefinition(
        name="docs.pdf_fill", version="1.0.0",
        description=("Fill a PDF form's fields and save a NEW filled copy in the Library (the original is "
                     "never modified). Use field names from docs.pdf_fields. Only use values the user gave or "
                     "that are in their memory/profile; never guess, and never fill signature fields."),
        input_schema={"type": "object", "properties": {
            "artifact_id": {"type": "string", "maxLength": 40},
            "values": {"type": "object", "additionalProperties": {"type": ["string", "boolean", "number"]}}},
            "required": ["artifact_id", "values"], "additionalProperties": False},
        output_schema={"type": "object"}, capabilities=["docs.write"], side_effect="local_write",
        idempotency="keyed", default_timeout_ms=30_000, execute=fill,
        display=lambda out: {"type": "document", "title": out.get("name", "Filled form"),
                             "artifact_id": out.get("artifact_id"), "subtitle": "Filled copy — review before sending",
                             "fields": (out.get("filled") or [])[:12], "mime": "application/pdf"},
    ))

    def create(ctx, a):
        rec = library.create(uid(ctx), title=a["title"], fmt=a["format"], content=a.get("content", ""),
                             columns=a.get("columns"), rows=a.get("rows"))
        return {"artifact_id": rec["artifact_id"], "name": rec["name"], "size": rec["size"]}

    registry.register(ToolDefinition(
        name="docs.create", version="1.0.0",
        description=("Create a document in the user's Library. format=md or pdf uses `content` (Markdown: "
                     "headings, lists, **bold**, tables); format=csv or xlsx uses `columns` + `rows`."),
        input_schema={"type": "object", "properties": {
            "title": {"type": "string", "maxLength": 120},
            "format": {"type": "string", "enum": ["md", "pdf", "csv", "xlsx"]},
            "content": {"type": "string", "maxLength": 60000},
            "columns": {"type": "array", "maxItems": 40, "items": {"type": "string", "maxLength": 80}},
            "rows": {"type": "array", "maxItems": 2000, "items": {"type": "array", "maxItems": 40}}},
            "required": ["title", "format"], "additionalProperties": False},
        output_schema={"type": "object"}, capabilities=["docs.write"], side_effect="local_write",
        idempotency="keyed", default_timeout_ms=90_000, execute=create,
        display=lambda out: {"type": "document", "title": out.get("name", "Document"),
                             "artifact_id": out.get("artifact_id"), "subtitle": "Created in your Library"},
    ))
